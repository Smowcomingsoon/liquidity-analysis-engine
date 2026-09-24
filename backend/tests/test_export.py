"""Il backtest esportato come foglio di calcolo.

Serve a rifare i conti fuori dall'app. Le due proprietà che lo rendono utile
invece che decorativo: i totali sono formule (cancelli tre righe e il riepilogo
segue) e le percentuali sono frazioni (4.7 scritto così con formato percentuale
diventerebbe 470%).
"""

from __future__ import annotations

import io

import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook

from app.analysis.backtest import run_backtest
from app.export import _ts, build_backtest_workbook

LORENZO = "Lorenzo (trend-follower)"


@pytest.fixture(scope="module")
def book():
    rng = np.random.default_rng(2)
    n = 600
    close = 100 * np.exp(np.cumsum(rng.normal(0.002, 0.02, n)))
    df = pd.DataFrame(
        {"open": close, "high": close * 1.01, "low": close * 0.99, "close": close,
         "volume": rng.uniform(1e6, 5e6, n)},
        index=pd.date_range("2022-01-01", periods=n, freq="D"),
    )
    bt = run_backtest(df, None, {}, strategy_name=LORENZO)
    blob = build_backtest_workbook(
        bt, symbol="kasusdt", strategy_name=LORENZO, chart_interval="1d",
        params={"len": 130, "average": "EMA"}, backtest_params={"initial_capital": 10000},
    )
    return bt, load_workbook(io.BytesIO(blob))


# --- il timestamp ---------------------------------------------------------


def test_seconds_and_milliseconds_both_land_in_this_century():
    """Il backtest usa entrambe le unità: trattarle tutte come secondi dava
    l'anno 53971 sulla data di inizio dati."""
    assert _ts(1_700_000_000).year == 2023
    assert _ts(1_700_000_000_000).year == 2023


def test_a_missing_timestamp_is_empty_not_1970():
    assert _ts(None) is None and _ts(0) is None


# --- la struttura ---------------------------------------------------------


def test_the_three_sheets_are_there(book):
    _bt, wb = book
    assert wb.sheetnames[:3] == ["Summary", "Trade", "Equity"]


def test_every_trade_has_a_row(book):
    bt, wb = book
    assert wb["Trade"].max_row >= bt.num_trades + 1


def test_the_charts_are_embedded(book):
    """«con tanto di grafico»: la curva del capitale, il calo dal massimo e il
    risultato per operazione."""
    _bt, wb = book
    assert len(wb["Equity"]._charts) == 2
    assert len(wb["Trade"]._charts) == 1


# --- le due proprietà che contano -----------------------------------------


def test_the_totals_are_formulas_not_pasted_numbers(book):
    bt, wb = book
    ws = wb["Trade"]
    total = ws.cell(row=bt.num_trades + 2, column=9).value
    assert isinstance(total, str) and total.startswith("=SUM(")


def test_the_summary_reads_from_the_trades_sheet(book):
    """Un riepilogo che resta fermo mentre le righe cambiano è peggio di uno
    assente, perché sembra ancora giusto."""
    _bt, wb = book
    ws = wb["Summary"]
    formulas = [ws.cell(r, 2).value for r in range(1, ws.max_row + 1)]
    text = " ".join(str(f) for f in formulas if isinstance(f, str))
    assert "Trade!" in text
    assert "COUNTIF(Trade!" in text


def test_percentages_are_fractions_with_a_percent_format(book):
    bt, wb = book
    if not bt.trades:
        pytest.skip("nessun trade su questa serie")
    cell = wb["Trade"].cell(row=2, column=8)
    assert abs(cell.value - bt.trades[0].pnl_pct / 100.0) < 1e-12
    assert "%" in cell.number_format


def test_the_drawdown_column_recomputes_the_peak(book):
    """È il conto che chiunque rifarebbe a mano per controllare il drawdown
    dichiarato: deve poterlo rifare il foglio."""
    _bt, wb = book
    assert str(wb["Equity"].cell(3, 3).value).startswith("=MAX(")
    assert str(wb["Equity"].cell(3, 4).value).startswith("=IF(")


def test_the_parameters_travel_with_the_file(book):
    _bt, wb = book
    ws = wb["Summary"]
    labels = [ws.cell(r, 1).value for r in range(1, ws.max_row + 1)]
    assert "len" in labels and "average" in labels


def test_an_empty_backtest_still_produces_a_readable_file():
    """Zero trade è un risultato, non un errore: il file deve aprirsi lo stesso."""
    from app.models import BacktestResult

    empty = BacktestResult(
        interval="1d", initial_capital=10000, final_equity=10000, total_return_pct=0.0,
        num_trades=0, win_rate_pct=None, max_drawdown_pct=0.0, profit_factor=None,
        equity_curve=[], trades=[],
    )
    wb = load_workbook(io.BytesIO(build_backtest_workbook(
        empty, symbol="x", strategy_name=LORENZO, chart_interval="1d",
        params={}, backtest_params={})))
    assert wb["Trade"].max_row == 1
